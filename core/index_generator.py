from typing import List

from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from models.figure import Figure, FigureStyle


class IndexGenerator:
    """
    Inserts a Word-native Table of Figures field:
        { TOC \\h \\z \\c "Figura" }

    This produces an updatable TOF identical to the one Word generates via
    Referencias → Insertar tabla de ilustraciones → Rótulo: Figura.

    After opening the document Word refreshes the field automatically
    (w:updateFields is set so it triggers on first open).
    The user can also right-click → "Actualizar campos" at any time.
    """

    def __init__(self, style: FigureStyle = None):
        self.style = style or FigureStyle()

    # ------------------------------------------------------------------
    def generate_index(self, document, figures: List[Figure]):
        # 1. Page break before the index
        pb = document.add_paragraph()
        pb.add_run().add_break(WD_BREAK.PAGE)

        # 2. Heading "Índice de Figuras"
        heading = document.add_paragraph()
        try:
            heading.style = document.styles["TOC Heading"]
        except KeyError:
            pass  # style not in document — plain formatting below is enough
        run = heading.add_run("Indice de Figuras")
        run.bold = True
        run.font.name = self.style.font_name
        run.font.size = Pt(self.style.font_size + 2)
        heading.paragraph_format.space_after = Pt(6)

        # 3. Paragraph that holds the TOF complex field
        tof_para = document.add_paragraph()
        self._insert_tof_field(tof_para._element, label="Figura")

        # 4. Tell Word to update all fields on open (covers SEQ numbers + TOF)
        self._set_update_fields_on_open(document)

    # ------------------------------------------------------------------
    @staticmethod
    def _insert_tof_field(para_element, label: str):
        """
        Insert { TOC \\h \\z \\c "label" } as a complex field.

        \\h  – entries become hyperlinks
        \\z  – hides tab/page numbers in web view
        \\c  – collect captions whose SEQ identifier matches <label>

        dirty="1" forces Word to recalculate the field on open.
        """
        # begin
        r_begin = OxmlElement("w:r")
        fc_begin = OxmlElement("w:fldChar")
        fc_begin.set(qn("w:fldCharType"), "begin")
        fc_begin.set(qn("w:dirty"), "1")
        r_begin.append(fc_begin)
        para_element.append(r_begin)

        # instruction
        r_instr = OxmlElement("w:r")
        instr = OxmlElement("w:instrText")
        instr.text = f' TOC \\h \\z \\c "{label}" '
        instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r_instr.append(instr)
        para_element.append(r_instr)

        # separate (no cached entries — Word fills them on update)
        r_sep = OxmlElement("w:r")
        fc_sep = OxmlElement("w:fldChar")
        fc_sep.set(qn("w:fldCharType"), "separate")
        r_sep.append(fc_sep)
        para_element.append(r_sep)

        # end
        r_end = OxmlElement("w:r")
        fc_end = OxmlElement("w:fldChar")
        fc_end.set(qn("w:fldCharType"), "end")
        r_end.append(fc_end)
        para_element.append(r_end)

    # ------------------------------------------------------------------
    @staticmethod
    def _set_update_fields_on_open(document):
        """
        Add <w:updateFields w:val="1"/> to document settings so Word
        automatically refreshes all fields (SEQ + TOF) when the file is opened.
        """
        settings_el = document.settings.element
        update_el = settings_el.find(qn("w:updateFields"))
        if update_el is None:
            update_el = OxmlElement("w:updateFields")
            settings_el.insert(0, update_el)
        update_el.set(qn("w:val"), "1")
